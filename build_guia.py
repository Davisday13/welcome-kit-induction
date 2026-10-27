# -*- coding: utf-8 -*-
"""Guia de entrega y exposicion, alineada con la rubrica del curso."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

NAVY = RGBColor(0x0E, 0x27, 0x3F)
ACCENT = RGBColor(0xB9, 0x77, 0x1B)
MUTED = RGBColor(0x55, 0x69, 0x7C)

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(4)
st.paragraph_format.line_spacing = 1.12

sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(0.85)
sec.top_margin = sec.bottom_margin = Inches(0.7)


def h(text, size=13, space=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = NAVY
    return p


def par(text, italic=False, color=None, size=10.5):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    return p


def bullet(text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    return p


t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Induction Project \u2014 Gu\u00eda de entrega")
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = NAVY

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_after = Pt(10)
r = sub.add_run("Davis Fern\u00e1ndez  \u00b7  IT Support  \u00b7  Octubre 2026  \u00b7  "
                "basada en la r\u00fabrica del curso")
r.font.size = Pt(9.5)
r.font.color.rgb = MUTED

h("1. Lo que ya est\u00e1 hecho")
bullet("Induction_Project_Davis_Fernandez_FINAL.pptx \u2014 4 diapositivas, mapa con la "
       "ruta y notas del orador de unas 535 palabras (4\u20135 minutos).")
bullet("Welcome_Kit_Escrito_Davis_Fernandez.docx \u2014 parte escrita de 338 palabras "
       "(el rango pedido es 250\u2013350).")
bullet("Esta gu\u00eda.")

h("2. Datos que faltan (m\u00e1ndamelos y lo dejo cerrado)")
bullet("Nombre completo tal como debe aparecer en la portada.")
bullet("Direcci\u00f3n exacta de tu base en La Riviera (calle y n\u00famero de casa).")
bullet("Confirmar la ruta: si giras en otro punto de la Calle 4 o entras por otra "
       "calle a PikTa, dime cu\u00e1l.")
bullet("D\u00f3nde estacionas cuando trabajas en el restaurante.")
bullet("Los dos lugares de comida reales (ahora son \u2018the bakery on the corner\u2019 y "
       "\u2018the market\u2019) y el precio aproximado.")
bullet("Ciudades que cubres, para confirmar David, Bugaba, Boquete y Volc\u00e1n.")
bullet("Extensi\u00f3n de emergencia o seguro m\u00e9dico, si quieres a\u00f1adirlo al 911.")
bullet("Foto tuya o logo (opcional, para la portada).")

h("3. C\u00f3mo cambiar el mapa por una captura de Google Maps")
par("El mapa dibujado ya cumple, pero una captura real suma en \u2018mapa incluido\u2019:")
bullet("Abre Google Maps, escribe tu direcci\u00f3n y activa la ruta desde tu casa.")
bullet("Captura la pantalla con la ruta en azul visible (Win + Shift + S).")
bullet("En PowerPoint: diapositiva 3, selecciona el mapa, Insertar > Im\u00e1genes y "
       "coloca la captura encima. Ajusta al mismo tama\u00f1o y env\u00eda el mapa dibujado "
       "al fondo o b\u00f3rralo.")
bullet("Deja la etiqueta de la ruta: salida, sem\u00e1foro donde giras y llegada.")

h("4. Chequeo contra la r\u00fabrica")
rows = [
    ("Contenido (25)", "Las 4 secciones est\u00e1n completas y claras", "Listo"),
    ("Gram\u00e1tica (25)", "Presente simple + continuo, pasado y futuro", "Listo"),
    ("Gram\u00e1tica (25)", "5+ preposiciones (in, on, at, by, with, next to\u2026)", "12 distintas"),
    ("Vocabulario (20)", "Direcciones, comida, primeros auxilios, profesi\u00f3n", "Listo"),
    ("Vocabulario (20)", "some / any / a lot of / many / much", "Las 5"),
    ("Presentaci\u00f3n (10)", "Mapa con la ruta marcada", "Listo"),
    ("Presentaci\u00f3n (10)", "4 diapositivas organizadas", "Listo"),
    ("Oral (20)", "4\u20135 minutos, sin leer", "Notas: ~535 palabras"),
    ("Extensi\u00f3n", "250\u2013350 palabras escritas", "345 en PPT / 338 en Word"),
]
tb = doc.add_table(rows=1, cols=3)
tb.style = "Light Grid Accent 1"
tb.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = tb.rows[0].cells
for i, txt in enumerate(["Criterio", "Qu\u00e9 pide", "Estado"]):
    hdr[i].text = txt
    for p in hdr[i].paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(9.5)
for a, b, c in rows:
    cells = tb.add_row().cells
    for i, txt in enumerate((a, b, c)):
        cells[i].text = txt
        for p in cells[i].paragraphs:
            for r in p.runs:
                r.font.size = Pt(9)

h("5. Guion oral (resumen en espa\u00f1ol)")
par("Diapositiva 1 \u2014 saludo y objetivo (30 s). Diapositiva 2 \u2014 perfil, que "
    "trabajas en toda la provincia, horario y rutina (1:10). Diapositiva 3 \u2014 "
    "tu base en La Riviera, la ruta a PikTa con First / Then / After that / "
    "Finally y el mapa (1:15). Diapositiva 4 \u2014 comida, la cena de ayer, el d\u00eda "
    "largo de ma\u00f1ana y seguridad (1:20). Cierre: \u201cThank you. Any questions?\u201d")
par("El texto completo en ingl\u00e9s est\u00e1 en las notas del orador de cada "
    "diapositiva (Vista > Notas).", italic=True, color=MUTED)

h("6. Fechas del proyecto")
bullet("Semana 3 \u2014 parte escrita en Word (borrador para revisi\u00f3n).")
bullet("Semana 4 \u2014 borrador en PPT u otra aplicaci\u00f3n.")
bullet("Semana 5 \u2014 versi\u00f3n final escrita en Classroom (PPT).")
bullet("Semana 6 \u2014 presentaci\u00f3n oral en la sesi\u00f3n en vivo.")

h("7. Antes de entregar")
bullet("Cerrar PowerPoint para poder actualizar el archivo con los datos finales.")
bullet("Reemplazar los tres campos entre corchetes y revisar que no quede ning\u00fan "
       "\u2018Maria Gonzalez\u2019.")
bullet("Exportar una copia en PDF (Archivo > Exportar > PDF).")
bullet("Practicar en voz alta dos veces con las notas: 4 a 5 minutos.")

out = r"C:\Users\DAVIS\Documents\Induction_Project\Guia_Induction_Project.docx"
doc.save(out)
print("Saved:", out)
