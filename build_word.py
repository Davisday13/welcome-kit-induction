# -*- coding: utf-8 -*-
"""Version escrita del Welcome Kit (250-350 palabras) para el borrador en Classroom."""

import re

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x0E, 0x27, 0x3F)
MUTED = RGBColor(0x55, 0x69, 0x7C)

SECTIONS = [
    ("1. About me",
     "My name is Davis Fern\u00e1ndez and I am an independent IT technician. I am "
     "friendly, organized and a little shy, but I enjoy solving problems and helping "
     "people. I do not work in one office all day: I visit clients all over Chiriqu\u00ed "
     "province, and my base is in La Riviera, in David. My main client is "
     "Restaurante PikTa, on Calle Ar\u00edstides Romero. I start work at 8:00 a.m. and I "
     "finish at 5:00 p.m., from Monday to Friday. This week I am setting up a "
     "network and I am visiting clients in Bugaba and Boquete."),
    ("2. How to get to my workplace",
     "My base is in Urbanizaci\u00f3n La Riviera, on Calle 5, but my main workplace is "
     "Restaurante PikTa, on Calle Ar\u00edstides Romero. From my base, go east on Calle 5 "
     "and turn left on Calle 4. Then go north to Calle Ar\u00edstides Romero and turn "
     "right. PikTa is one block ahead, near Ozaga Market, and there is parking "
     "behind the building. The map with my route is on slide 3, and it takes about "
     "five minutes by car."),
    ("3. Around the office",
     "There are a lot of places to eat near my office. The bakery on the corner is "
     "cheap: coffee, bread and fruit for about two dollars. At the market, a typical "
     "lunch is rice with chicken, beans and plantain for about five dollars. Some "
     "places are expensive and many people eat at the market, but I don't have much "
     "time, so I usually eat near my office. Yesterday I finished work and I went to "
     "dinner with my family. Tomorrow I will have a long day: I will visit three "
     "clients in David, Bugaba and Boquete."),
    ("4. Safety first",
     "The first-aid kit is in my car, in a red box, and I have the key. If someone "
     "is bleeding, apply pressure and call 911. The emergency number in Panama is "
     "911. When I travel, I check the car before a long trip, I share my route with "
     "my family, and I back up my data before every visit."),
]

text = " ".join(body for _, body in SECTIONS)
count = len(re.findall(r"[A-Za-z0-9']+", text))

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15

sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(1.0)
sec.top_margin = sec.bottom_margin = Inches(0.9)

t = doc.add_paragraph()
t.paragraph_format.space_after = Pt(2)
r = t.add_run("Welcome Kit for a New Colleague")
r.bold = True
r.font.size = Pt(22)
r.font.color.rgb = NAVY

who = doc.add_paragraph()
who.paragraph_format.space_after = Pt(2)
r = who.add_run("Davis Fern\u00e1ndez  \u00b7  Independent IT Technician  \u00b7  "
                "David, Chiriqu\u00ed  \u00b7  October 2026")
r.font.size = Pt(9.5)
r.font.color.rgb = MUTED

wc = doc.add_paragraph()
wc.paragraph_format.space_after = Pt(14)
r = wc.add_run(f"Written part ({count} words, map on slide 3)")
r.italic = True
r.font.size = Pt(9.5)
r.font.color.rgb = MUTED

for title, body in SECTIONS:
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(2)
    hr = h.add_run(title)
    hr.bold = True
    hr.font.size = Pt(13)
    hr.font.color.rgb = NAVY
    p = doc.add_paragraph(body)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

out = r"C:\Users\DAVIS\Documents\Induction_Project\Welcome_Kit_Escrito_Davis_Fernandez.docx"
doc.save(out)
print("Saved:", out)
print("Words in written part:", count)
